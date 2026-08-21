"""
# mypy: ignore-errors
infrastructure/local_storage/excel_workbook_store.py — ExcelSession,
the pandas/openpyxl-backed in-memory/on-disk workbook
AI Model Coder CLI v1.51.0 (Clean Architecture refactor, Phase C, Context #4)

Extracted 2026-08-18 from claude_excel.py's ExcelSession class,
unmodified in behavior. Kept as a single class rather than split
method-by-method between domain/ and infrastructure/ — same reasoning
as PptxSession in infrastructure/local_storage/pptx_deck_store.py:
__init__/_load/save/_write_charts depend on pandas/openpyxl and touch
disk, while summary()/undo()/_snapshot()/apply_code()'s safety check
are pure logic, but they all operate on the same `self.sheets` state.
"""

from domain.excel import _DENYLIST

try:
    import pandas as pd
except ImportError:
    pd = None


class ExcelSession:
    def __init__(self, input_path=None, sheet_name=None):
        if pd is None:
            raise ImportError("pandas is required for --excel (pip install pandas openpyxl)")
        self.sheets = {}
        self._history_stack = []  # for /undo — list of {name: df.copy()} snapshots
        self._pending_charts = []  # (sheet, chart_type, title, categories_col, value_cols)

        if input_path:
            self._load(input_path, sheet_name)
        else:
            self.sheets["Sheet1"] = pd.DataFrame()

    def _load(self, path, sheet_name=None):
        if path.lower().endswith(".csv"):
            self.sheets["Sheet1"] = pd.read_csv(path)
        else:
            all_sheets = pd.read_excel(path, sheet_name=None)
            if sheet_name:
                if sheet_name not in all_sheets:
                    raise ValueError(f"Sheet {sheet_name!r} not found; have {list(all_sheets)}")
                self.sheets = {sheet_name: all_sheets[sheet_name]}
            else:
                self.sheets = all_sheets

    # ── context for the model ───────────────────────────────────────────

    def summary(self):
        parts = []
        for name, df in self.sheets.items():
            cols = ", ".join(f"{c} ({df[c].dtype})" for c in df.columns[:30])
            parts.append(
                f"Sheet {name!r}: {df.shape[0]} rows x {df.shape[1]} cols. " f"Columns: {cols or '(empty)'}"
            )
            if not df.empty:
                parts.append(f"First rows of {name!r}:\n{df.head(5).to_string()}")
        return "\n".join(parts) if parts else "(no data loaded yet)"

    # ── applying a model turn ───────────────────────────────────────────

    def _snapshot(self):
        self._history_stack.append({k: v.copy() for k, v in self.sheets.items()})
        if len(self._history_stack) > 20:
            self._history_stack.pop(0)

    def undo(self):
        if not self._history_stack:
            return False
        self.sheets = self._history_stack.pop()
        return True

    def apply_code(self, code):
        """Run model-generated code against `self.sheets`. Returns (ok, message)."""
        lowered = code.lower()
        for bad in _DENYLIST:
            if bad in lowered:
                return False, f"[blocked] generated code used a disallowed construct: {bad!r}"

        self._snapshot()
        local_ns = {
            "sheets": self.sheets,
            "pd": pd,
            "add_chart": self._add_chart,
        }
        try:
            exec(
                compile(code, "<excel-turn>", "exec"),
                {
                    "__builtins__": {
                        "len": len,
                        "range": range,
                        "sum": sum,
                        "min": min,
                        "max": max,
                        "round": round,
                        "sorted": sorted,
                        "list": list,
                        "dict": dict,
                        "str": str,
                        "int": int,
                        "float": float,
                        "bool": bool,
                        "enumerate": enumerate,
                        "zip": zip,
                        "abs": abs,
                    }
                },
                local_ns,
            )
        except Exception as e:
            self.undo()
            return False, f"[ERROR] generated code failed: {e}"
        self.sheets = local_ns["sheets"]
        return True, "applied"

    def _add_chart(self, sheet, chart_type, title, categories_col, value_cols):
        self._pending_charts.append((sheet, chart_type, title, categories_col, value_cols))

    # ── persistence ──────────────────────────────────────────────────────

    def save(self, output_path):
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for name, df in self.sheets.items():
                df.to_excel(writer, sheet_name=name[:31], index=False)

        if self._pending_charts:
            self._write_charts(output_path)

    def _write_charts(self, output_path):
        import openpyxl
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference

        wb = openpyxl.load_workbook(output_path)
        for sheet, chart_type, title, categories_col, value_cols in self._pending_charts:
            sheet_key = sheet[:31]
            if sheet_key not in wb.sheetnames or sheet_key not in self.sheets:
                continue
            ws = wb[sheet_key]
            df = self.sheets[sheet_key]
            if categories_col not in df.columns:
                continue
            cat_idx = df.columns.get_loc(categories_col) + 1
            n_rows = df.shape[0]

            chart_cls = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_type, BarChart)
            chart = chart_cls()
            chart.title = title or f"{sheet_key} chart"
            cats = Reference(ws, min_col=cat_idx, min_row=2, max_row=n_rows + 1)

            for col_name in value_cols:
                if col_name not in df.columns:
                    continue
                val_idx = df.columns.get_loc(col_name) + 1
                data = Reference(ws, min_col=val_idx, min_row=1, max_row=n_rows + 1)
                chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            anchor_col = chr(ord("A") + df.shape[1] + 2)
            ws.add_chart(chart, f"{anchor_col}2")
        wb.save(output_path)
        self._pending_charts = []
